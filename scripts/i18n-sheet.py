#!/usr/bin/env python3
"""Exchange Chronicle's Android and web translations with a human-filled XLSX."""

import argparse
from copy import deepcopy
import json
from pathlib import Path
import re
import sys
import xml.etree.ElementTree as ET

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Alignment, Font


ANDROID = [
    (surface, Path(f"chronicle/app/src/{flavor}/res/values/{name}.xml"))
    for surface, flavor in (("android-play", "main"), ("android-research", "googleServices"))
    for name in ("strings", "consent_copy")
]
WEB = Path("chronicle-web/src/modern/i18n/en/translation.json")
HEADERS = ("surface", "file", "key", "english", "spanish", "notes")
EXCLUDED = {"user_target_child", "user_other", "user_unassigned"}
TOKENS = re.compile(r"%%|%(?:\d+\$)?[sd]|\{\{[^{}]+\}\}")


def placeholders(text):
    return set(TOKENS.findall(text)) - {"%%"}


def android_text(element):
    return re.sub(r"\\(['\"])", r"\1", "".join(element.itertext()))


def resources(tree):
    return [element for element in tree
            if element.tag in {"string", "string-array", "plurals"}
            and element.get("translatable") != "false"
            and element.get("name") not in EXCLUDED]


def items(element):
    name = element.attrib["name"]
    if element.tag == "string":
        return [(name, element)]
    return [(f"{name}[{item.get('quantity') if element.tag == 'plurals' else index}]", item)
            for index, item in enumerate(element.findall("item"))]


def flatten(value, path=()):
    if isinstance(value, dict):
        for key, child in value.items():
            yield from flatten(child, (*path, key))
    elif isinstance(value, list):
        for index, child in enumerate(value):
            yield from flatten(child, (*path, index))
    elif isinstance(value, str):
        yield ".".join(map(str, path)), value, path


def target_path(path, lang):
    if path == WEB:
        return WEB.parent.parent / lang / WEB.name
    return path.parent.parent / f"values-{lang}" / path.name


def read_json(path):
    return json.loads(path.read_text(encoding="utf-8"))


def catalog(root, lang):
    rows = {}
    trees = {}
    for surface, path in ANDROID:
        tree = ET.parse(root / path).getroot()
        trees[path] = tree
        translated = root / target_path(path, lang)
        existing = {}
        if translated.exists():
            existing = {key: android_text(item)
                        for element in resources(ET.parse(translated).getroot())
                        for key, item in items(element)}
        for element in resources(tree):
            for key, item in items(element):
                if item.get("translatable") != "false":
                    rows[(surface, path.as_posix(), key)] = (android_text(item), existing.get(key, ""))
    english = read_json(root / WEB)
    translated = root / target_path(WEB, lang)
    existing = {key: value for key, value, _ in flatten(read_json(translated))} if translated.exists() else {}
    for key, value, _ in flatten(english):
        rows[("web", WEB.as_posix(), key)] = (value, existing.get(key, ""))
    return rows, trees, english


def export_sheet(root, lang, output):
    rows, _, _ = catalog(root, lang)
    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "Strings"
    sheet.append(HEADERS)
    sheet.freeze_panes = "A2"
    for identity, (english, spanish) in rows.items():
        notes = sorted(placeholders(english))
        if r"\n" in english:
            notes.append(r"\n")
        notes.extend(re.findall(r"\$t\([^)]*\)", english))
        sheet.append((*identity, english, spanish, ", ".join(notes)))
    for row in sheet:
        for cell in row:
            cell.data_type = "s"  # Strings beginning with '=' must stay text.
            cell.alignment = Alignment(wrap_text=True, vertical="top")
            if cell.row == 1:
                cell.font = Font(bold=True)
    for column, width in zip("ABCDEF", (22, 58, 48, 80, 80, 32)):
        sheet.column_dimensions[column].width = width
    sheet.auto_filter.ref = sheet.dimensions
    output.parent.mkdir(parents=True, exist_ok=True)
    workbook.save(output)
    workbook.close()
    for surface in ("android-play", "android-research", "web"):
        print(f"{surface}: {sum(identity[0] == surface for identity in rows)} rows")
    print(f"Exported {len(rows)} rows to {output}")


def read_sheet(source, rows):
    workbook = load_workbook(source, read_only=True, data_only=False)
    translations = {}
    rejected = []
    try:
        if "Strings" not in workbook.sheetnames:
            raise ValueError("Workbook has no Strings sheet")
        sheet = workbook["Strings"]
        if tuple(cell.value for cell in next(sheet.iter_rows(max_row=1, max_col=len(HEADERS)))) != HEADERS:
            raise ValueError(f"Expected columns: {', '.join(HEADERS)}")
        seen = set()
        for number, cells in enumerate(sheet.iter_rows(min_row=2, max_col=6), 2):
            surface, file, key, english, spanish, _ = (cell.value for cell in cells)
            if spanish is None or spanish == "":
                continue
            identity = (surface, file, key)
            reason = None
            if identity not in rows:
                reason = "unknown surface/file/key"
            elif identity in seen:
                reason = "duplicate row"
            elif english != rows[identity][0]:
                reason = "English differs from the current source; export again"
            elif not isinstance(spanish, str) or cells[4].data_type == "f":
                reason = "spanish must be text, not a number or formula"
            elif placeholders(spanish) != placeholders(english):
                reason = (f"placeholder mismatch: expected {sorted(placeholders(english))}, "
                          f"got {sorted(placeholders(spanish))}")
            elif r"\n" in english and r"\n" not in spanish:
                reason = r"missing literal \n"
            seen.add(identity)
            if reason:
                rejected.append(f"Row {number} ({key}) rejected: {reason}")
            else:
                translations[identity] = spanish
    finally:
        workbook.close()
    if rejected:
        raise ValueError("\n".join(rejected) + "\nImport aborted; no files written.")
    return translations


def merge_web(existing, english, translations, path=()):
    """Keep existing data, including JSON array positions."""
    if path in translations:
        return translations[path]
    if not any(key[:len(path)] == path for key in translations):
        return existing
    if isinstance(english, dict):
        result = dict(existing) if isinstance(existing, dict) else {}
        for key, child in english.items():
            child_path = (*path, key)
            if any(target[:len(child_path)] == child_path for target in translations):
                result[key] = merge_web(result.get(key), child, translations, child_path)
        return result
    if isinstance(english, list):
        result = list(existing) if isinstance(existing, list) else []
        for index, child in enumerate(english):
            child_path = (*path, index)
            if any(target[:len(child_path)] == child_path for target in translations):
                # Missing earlier indices fall back to English, keeping array positions.
                while len(result) <= index:
                    result.append(deepcopy(english[len(result)]))
                result[index] = merge_web(result[index], child, translations, child_path)
        return result
    return existing


def english_order(value, english):
    if isinstance(value, dict) and isinstance(english, dict):
        keys = [key for key in english if key in value]
        keys.extend(key for key in value if key not in english)
        return {key: english_order(value[key], english.get(key)) for key in keys}
    if isinstance(value, list) and isinstance(english, list):
        return [english_order(child, english[index] if index < len(english) else None)
                for index, child in enumerate(value)]
    return value


def import_sheet(root, lang, source):
    rows, trees, english = catalog(root, lang)
    translations = read_sheet(source, rows)
    outputs = {}
    for surface, path in ANDROID:
        output = ET.Element("resources")
        for element in resources(trees[path]):
            entries = items(element)
            values = [translations.get((surface, path.as_posix(), key)) for key, _ in entries]
            if not entries or any(value is None for value in values):
                if element.tag != "string":
                    print(f"Warning: omitted incomplete {element.tag} {path}:{element.get('name')}", file=sys.stderr)
                continue
            translated = ET.SubElement(output, element.tag, element.attrib)
            for (_, original), value in zip(entries, values):
                item = translated if element.tag == "string" else ET.SubElement(translated, "item", original.attrib)
                item.text = value.replace("'", r"\'").replace('"', r'\"')
        ET.indent(output, space="    ")
        outputs[root / target_path(path, lang)] = ET.tostring(output, encoding="unicode", xml_declaration=True) + "\n"
    web_path = root / target_path(WEB, lang)
    existing = read_json(web_path) if web_path.exists() else {}
    web_translations = {path: translations[("web", WEB.as_posix(), key)]
                        for key, _, path in flatten(english)
                        if ("web", WEB.as_posix(), key) in translations}
    merged = english_order(merge_web(existing, english, web_translations), english)
    outputs[web_path] = json.dumps(merged, ensure_ascii=False, indent=2) + "\n"
    for path, content in outputs.items():
        path.parent.mkdir(parents=True, exist_ok=True)
        path.write_text(content, encoding="utf-8")
        print(f"Wrote {path}")
    print(f"Imported {len(translations)} filled rows (incomplete Android groups omitted).")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    default_root = Path(__file__).resolve().parent.parent
    parser.add_argument("--root", type=Path, default=default_root)
    commands = parser.add_subparsers(dest="command", required=True)
    for command, flag in (("export", "--out"), ("import", "--in")):
        sub = commands.add_parser(command)
        sub.add_argument("--lang", required=True)
        sub.add_argument("--root", type=Path, default=argparse.SUPPRESS)
        sub.add_argument(flag, dest="sheet", type=Path, required=True)
    args = parser.parse_args()
    if not re.fullmatch(r"[a-z]{2,3}(?:-r[A-Z]{2})?", args.lang) or args.lang == "en":
        parser.error("--lang must be a non-English Android language qualifier, e.g. es or pt-rBR")
    try:
        if args.command == "export":
            export_sheet(args.root, args.lang, args.sheet)
        else:
            import_sheet(args.root, args.lang, args.sheet)
    except (OSError, ValueError, ET.ParseError) as error:
        parser.exit(1, f"{error}\n")


if __name__ == "__main__":
    main()
